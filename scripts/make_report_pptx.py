"""Build the W2 report deck: real YOLOv8 backbone for the dual-stream skeleton.

Run with the bundled office runtime (python-pptx):
    <bundled-python> scripts/make_report_pptx.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "experiments" / "w2_yolov8_backbone" / "figures"
OUT = ROOT / "experiments" / "w2_yolov8_backbone" / "双流骨架YOLOv8接入汇报.pptx"

INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
RED = RGBColor(0xDC, 0x26, 0x26)
GREEN = RGBColor(0x15, 0x80, 0x3D)
AMBER = RGBColor(0xB4, 0x53, 0x09)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x0F, 0x17, 0x2A)
FONT = "Microsoft YaHei"
MONO = "Consolas"

W, H = 13.333, 7.5


def add_text(slide, x, y, w, h, text, size=14, color=INK, bold=False, font=FONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15, space_after=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.shadow.inherit = False
    # Drop the theme style reference: LibreOffice/PowerPoint otherwise re-apply
    # the theme's drop shadow even when effectLst is empty.
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.2)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def slide_header(slide, kicker, title, accent=BLUE):
    add_rect(slide, 0.55, 0.52, 0.09, 0.42, accent)
    add_text(slide, 0.78, 0.5, 9.5, 0.3, kicker, 11, accent, bold=True)
    add_text(slide, 0.78, 0.78, 11.8, 0.55, title, 26, INK, bold=True)
    add_rect(slide, 0.55, 1.45, 12.23, 0.012, RGBColor(0xE5, 0xE7, 0xEB))


def slide_footer(slide, page):
    add_text(slide, 0.55, 7.08, 8.5, 0.25, "天巡双瞳 · tianxun-dual-patrol · W2 双流骨架", 9, MUTED)
    add_text(slide, 11.6, 7.08, 1.2, 0.25, f"{page:02d}", 9, MUTED, align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, items, size=14, gap=0.52, color=INK):
    yy = y
    for head, body in items:
        add_rect(slide, x, yy + 0.07, 0.09, 0.09, BLUE)
        add_text(slide, x + 0.22, yy - 0.02, w - 0.22, 0.32, head, size, color, bold=True)
        if body:
            add_text(slide, x + 0.22, yy + 0.28, w - 0.22, 0.5, body, size - 3.5, MUTED,
                     line_spacing=1.1)
            yy += gap + 0.24
        else:
            yy += gap


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    page = 0

    # ---------------------------------------------------------------- cover
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, DARK)
    add_rect(s, 0, 0, 0.18, H, BLUE)
    add_rect(s, 10.4, 0, 2.93, H, RGBColor(0x17, 0x24, 0x3D))
    add_text(s, 0.95, 1.55, 8.6, 0.4, "W2 · 双流骨架里程碑", 14, RGBColor(0x93, 0xC5, 0xFD), bold=True)
    add_text(s, 0.95, 2.0, 8.9, 1.5,
             "双流骨架接入真实 YOLOv8\nRGB / 红外 两路独立 Backbone", 36, WHITE, bold=True, line_spacing=1.18)
    add_text(s, 0.95, 3.75, 8.4, 1.0,
             "用 yolov8n-obb 前 10 层替换 _DummyBackbone，输出 P2 / P3 / P4 / P5，\n"
             "在随机数据上完成单路与整网前向验证。", 15, RGBColor(0xCB, 0xD5, 0xE1), line_spacing=1.4)
    add_rect(s, 0.95, 5.0, 1.1, 0.05, BLUE)
    add_text(s, 0.95, 5.2, 8.0, 0.35, "2026-10-27  ·  张涵来 / 林宸宇  ·  组内技术汇报", 12,
             RGBColor(0x94, 0xA3, 0xB8))
    add_text(s, 10.95, 5.2, 1.95, 1.2, "P2 160×160\nP3  80×80\nP4  40×40\nP5  20×20", 11,
             RGBColor(0x93, 0xC5, 0xFD), line_spacing=1.35)
    page += 1

    # ------------------------------------------------------------- objective
    s = prs.slides.add_slide(blank)
    slide_header(s, "01 / 目标", "把假 backbone 换成真 YOLOv8，且两路互不共享")
    cards = [
        ("替换", "_DummyBackbone → YOLOv8", "加载 yolov8n-obb-initial.pt，取前 10 层（0–9，含 SPPF）。", BLUE),
        ("独立", "RGB / 红外 两套实例", "各自从同一权重初始化，参数零共享；红外默认复制为 3 通道。", RED),
        ("对齐", "P2–P5 四层特征", "层号 2 / 4 / 6 / 9，stride 4 / 8 / 16 / 32，通道 32 / 64 / 128 / 256。", GREEN),
    ]
    x = 0.55
    for tag, head, body, accent in cards:
        add_rect(s, x, 1.85, 3.95, 3.05, LIGHT, radius=True)
        add_rect(s, x, 1.85, 3.95, 0.08, accent)
        add_text(s, x + 0.3, 2.12, 1.2, 0.3, tag, 12, accent, bold=True)
        add_text(s, x + 0.3, 2.48, 3.35, 0.75, head, 18, INK, bold=True, line_spacing=1.15)
        add_text(s, x + 0.3, 3.3, 3.35, 1.3, body, 12, MUTED, line_spacing=1.35)
        x += 4.14
    add_text(s, 0.55, 5.25, 12.2, 0.9,
             "验收标准：单路输入 640×640 时 P2 (1,32,160,160)、P3 (1,64,80,80)、P4 (1,128,40,40)、"
             "P5 (1,256,20,20)；整网 6 通道随机数据前向不报错。", 13, INK, line_spacing=1.4)
    slide_footer(s, page)
    page += 1

    # ----------------------------------------------------------- environment
    s = prs.slides.add_slide(blank)
    slide_header(s, "02 / 环境", "editable 安装 ultralytics 8.4.178，本机仅 CPU")
    rows = [
        ("Python", "3.14.5", "项目内 .venv"),
        ("ultralytics", "8.4.178（editable）", "pip install -e ./ultralytics"),
        ("PyTorch", "2.14.1+cpu", "PyPI cp314 Windows wheel"),
        ("CUDA", "不可用", "本机只有 Intel UHD，无 NVIDIA 驱动"),
        ("权重", "yolov8n-obb-initial.pt", "6.3 MB，fp16 保存 → 载入后转 fp32"),
    ]
    y = 1.78
    for i, (k, v, note) in enumerate(rows):
        add_rect(s, 0.55, y, 12.23, 0.62, LIGHT if i % 2 == 0 else WHITE)
        add_text(s, 0.8, y + 0.16, 2.0, 0.3, k, 13, MUTED, bold=True)
        add_text(s, 3.0, y + 0.16, 4.3, 0.3, v, 13, INK, bold=True)
        add_text(s, 7.4, y + 0.18, 5.1, 0.3, note, 11, MUTED)
        y += 0.66
    add_rect(s, 0.55, 5.3, 12.23, 1.15, RGBColor(0xFE, 0xF2, 0xF2), RGBColor(0xFE, 0xCA, 0xCA), radius=True)
    add_text(s, 0.85, 5.48, 11.7, 0.8,
             "CUDA 结论：本机无 NVIDIA GPU，torch.cuda.is_available() = False。本汇报所有形状与耗时均为 CPU 实测；\n"
             "代码不绑定设备，迁到 GPU 机器后把 .to(device) 打开即可复跑同样的自测脚本。", 11.5,
             RGBColor(0x99, 0x1B, 0x1B), line_spacing=1.45)
    slide_footer(s, page)
    page += 1

    # ------------------------------------------------------------ schematic
    s = prs.slides.add_slide(blank)
    slide_header(s, "03 / 骨架", "两条独立流水线，在 QAFM 处按层融合")
    img = FIG / "architecture.png"
    pic_w = 9.5
    pic_x = (W - pic_w) / 2
    pic = s.shapes.add_picture(str(img), Inches(pic_x), Inches(1.62), width=Inches(pic_w))
    pic_h = pic.height / 914400
    add_text(s, pic_x, 1.62 + pic_h + 0.1, pic_w, 0.3,
             "图 1  RGB 与红外各一条 YOLOv8 主干（0–9 层）；P2–P5 特征形状两路一致，QAFM 逐层融合后送入检测头。",
             10, MUTED)
    slide_footer(s, page)
    page += 1

    # --------------------------------------------------------- implementation
    s = prs.slides.add_slide(blank)
    slide_header(s, "04 / 实现", "新增 yolov8_backbone.py，保持原 {level: feature} 契约")
    bullets(s, 0.55, 1.85, 5.9, [
        ("加载与裁剪", "torch.load(weights_only=False) → model.model[:10]；导出权重为 fp16，统一 .float() 后再训练。"),
        ("特征抽头", "顺序执行 10 层，在 index 2 / 4 / 6 / 9 处收集 P2–P5，不依赖 forward hook。"),
        ("单通道可选", "in_ch=1 时把首层 3×3 卷积按输入通道求和折叠，默认仍走 3 通道复制。"),
        ("双实例", "DualStreamBackbone 构造两个独立 YOLOv8Backbone，并断言不是同一对象。"),
    ], size=13.5, gap=0.42)
    add_rect(s, 6.75, 1.85, 6.0, 3.9, DARK, radius=True)
    add_text(s, 7.05, 2.05, 5.4, 0.3, "src/models/yolov8_backbone.py", 10, RGBColor(0x93, 0xC5, 0xFD), font=MONO)
    code = (
        "bb = YOLOv8Backbone(weights, in_ch=3)\n"
        "feats = bb(x)\n"
        "# {2: P2, 3: P3, 4: P4, 5: P5}\n"
        "\n"
        "self.rgb = YOLOv8Backbone(weights, 3)\n"
        "self.thermal = YOLOv8Backbone(\n"
        "    weights, thermal_in_ch)\n"
        "assert self.rgb is not self.thermal\n"
        "\n"
        "# inferred from a real forward pass\n"
        "# {2: 32, 3: 64, 4: 128, 5: 256}"
    )
    add_text(s, 7.05, 2.45, 5.4, 3.1, code, 11.5, RGBColor(0xE2, 0xE8, 0xF0), font=MONO, line_spacing=1.32)
    add_text(s, 0.55, 5.95, 12.2, 0.9,
             "dual_stream_yolo.py：DualStreamDetector 新增 weights 参数，默认指向 weights/baseline_obb/yolov8n-obb-initial.pt；"
             "backbone 通道数由真实前向推断，不再硬编码。", 12, MUTED, line_spacing=1.4)
    slide_footer(s, page)
    page += 1

    # ------------------------------------------------------------ shape table
    s = prs.slides.add_slide(blank)
    slide_header(s, "05 / 单路自测", "P2–P5 形状与期望逐项一致")
    data = [
        ["层级", "层号", "stride", "通道", "实测形状（batch=1）", "期望", "结果"],
        ["P2", "2", "4", "32", "(1, 32, 160, 160)", "(1, 32, 160, 160)", "通过"],
        ["P3", "4", "8", "64", "(1, 64, 80, 80)", "(1, 64, 80, 80)", "通过"],
        ["P4", "6", "16", "128", "(1, 128, 40, 40)", "(1, 128, 40, 40)", "通过"],
        ["P5", "9", "32", "256", "(1, 256, 20, 20)", "(1, 256, 20, 20)", "通过"],
    ]
    rows, cols = len(data), len(data[0])
    shape = s.shapes.add_table(rows, cols, Inches(0.55), Inches(1.8), Inches(12.23), Inches(2.55))
    table = shape.table
    widths = [0.95, 0.85, 0.95, 0.95, 3.0, 3.0, 0.95]
    for i, wd in enumerate(widths):
        table.columns[i].width = Inches(wd)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.1)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.runs[0]
            run.font.size = Pt(12 if r else 11.5)
            run.font.bold = r == 0 or c == 0
            run.font.name = FONT
            run.font.color.rgb = WHITE if r == 0 else (GREEN if c == 6 else INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (LIGHT if r % 2 else WHITE)
    add_text(s, 0.55, 4.75, 12.2, 0.4, "同一套检查还覆盖 batch=2、3 通道红外、1 通道红外（首层卷积已折叠）。", 12, MUTED)
    stats = [("1,272,656", "单路 backbone 参数量"), ("0", "两路共享张量"), ("2,735,494", "整网参数量（5 类）")]
    x = 0.55
    for num, label in stats:
        add_rect(s, x, 5.35, 3.95, 1.15, LIGHT, radius=True)
        add_text(s, x + 0.3, 5.52, 3.4, 0.45, num, 22, BLUE, bold=True)
        add_text(s, x + 0.3, 6.02, 3.4, 0.3, label, 11, MUTED)
        x += 4.14
    slide_footer(s, page)
    page += 1

    # -------------------------------------------------------- single screenshot
    s = prs.slides.add_slide(blank)
    slide_header(s, "06 / 单路截图", "scripts/test_yolov8_backbone.py 实测输出")
    pic = s.shapes.add_picture(str(FIG / "backbone_selftest.png"), Inches(0.55), Inches(1.65),
                               height=Inches(5.1))
    add_text(s, 0.55 + pic.width / 914400 + 0.45, 1.75, 5.15, 0.35, "检查项", 15, INK, bold=True)
    items = [
        ("RGB 3 通道", "640×640，形状全部命中"),
        ("batch=2", "批维正确扩展"),
        ("红外 3 通道", "复制通道直接复用 3 通道 stem"),
        ("红外 1 通道", "首层卷积折叠为 1 输入通道，形状不变"),
        ("独立实例", "共享张量 0，参数不共享"),
    ]
    y = 2.2
    for head, body in items:
        add_text(s, 0.55 + pic.width / 914400 + 0.45, y, 5.1, 0.3, "✓  " + head, 12.5, GREEN, bold=True)
        add_text(s, 0.55 + pic.width / 914400 + 0.75, y + 0.26, 4.8, 0.3, body, 11, MUTED)
        y += 0.78
    slide_footer(s, page)
    page += 1

    # ------------------------------------------------------ end-to-end screenshot
    s = prs.slides.add_slide(blank)
    slide_header(s, "07 / 整网前向", "随机数据跑通双流检测器（P2 已启用）")
    s.shapes.add_picture(str(FIG / "dual_forward.png"), Inches(0.55), Inches(1.68), height=Inches(4.95))
    add_text(s, 6.75, 1.78, 5.9, 0.35, "运行配置", 15, INK, bold=True)
    cfg = [
        ("输入", "torch.randn(1, 6, 640, 640)，前 3 通道 RGB、后 3 通道红外"),
        ("模型", "DualStreamDetector(num_classes=5, use_tiny_p2=True)"),
        ("输出", "P2–P5 检测头 9×160×160 … 9×20×20；gate 每层 (B,2)"),
        ("模式", "eval 两种入参 + train（模态 dropout、辅助头、光照头）"),
        ("耗时", "CPU 单次前向约 108–164 ms"),
    ]
    y = 2.22
    for k, v in cfg:
        add_text(s, 6.75, y, 1.0, 0.3, k, 12, MUTED, bold=True)
        add_text(s, 7.5, y, 5.1, 0.55, v, 12, INK, line_spacing=1.25)
        y += 0.72
    add_rect(s, 6.75, 5.95, 5.9, 0.62, RGBColor(0xF0, 0xFD, 0xF4), RGBColor(0xBB, 0xF7, 0xD0), radius=True)
    add_text(s, 6.95, 6.12, 5.5, 0.3, "ALL END-TO-END CHECKS PASSED", 11, GREEN, bold=True, font=MONO)
    slide_footer(s, page)
    page += 1

    # ------------------------------------------------------------ conclusion
    s = prs.slides.add_slide(blank)
    slide_header(s, "08 / 结论", "骨架已换真，下一步是 neck / Detect 与真数据")
    add_text(s, 0.55, 1.8, 6.0, 0.35, "已完成", 15, GREEN, bold=True)
    done = [
        "ultralytics 8.4.178 editable 安装，版本与 CUDA 状态已核实",
        "yolov8_backbone.py：10 层裁剪、P2–P5 抽头、fp16→fp32",
        "红外 3 通道默认打通，1 通道折叠作为可选开关",
        "单路形状、双实例独立性、整网前后向全部自测通过",
    ]
    y = 2.22
    for t in done:
        add_text(s, 0.55, y, 6.05, 0.32, "•  " + t, 11.5, INK)
        y += 0.42
    add_text(s, 7.0, 1.8, 5.8, 0.35, "下一步", 15, AMBER, bold=True)
    nxt = [
        "接真实 neck + YOLOv8 Detect（P2 分支、OBB 角度）",
        "用 DroneVehicle 成对数据跑通训练与对齐检查",
        "有 NVIDIA 机器后补 CUDA 复跑与显存/吞吐数据",
        "QAFM 门控在真实模态质量下的可视化",
    ]
    y = 2.22
    for t in nxt:
        add_text(s, 7.0, y, 5.8, 0.32, "•  " + t, 11.5, INK)
        y += 0.42
    add_rect(s, 0.55, 4.95, 12.23, 1.55, LIGHT, radius=True)
    add_text(s, 0.85, 5.12, 11.6, 0.3, "复跑命令", 11, MUTED, bold=True)
    add_text(s, 0.85, 5.42, 11.6, 1.0,
             ".venv\\Scripts\\python.exe scripts\\test_yolov8_backbone.py\n"
             ".venv\\Scripts\\python.exe scripts\\test_dual_stream_forward.py",
             11.5, INK, font=MONO, line_spacing=1.5)
    slide_footer(s, page)
    page += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
